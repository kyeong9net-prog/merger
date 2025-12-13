"""Test helper functions for creating test Excel files (Phase 5)"""
import os
import tempfile
from typing import List, Any, Optional
from openpyxl import Workbook
from datetime import datetime


def create_test_excel_file(
    filepath: str,
    sheet_name: str = "summary",
    header: Optional[List[Any]] = None,
    data_row: Optional[List[Any]] = None,
    has_merged_cells: bool = False,
    has_formula: bool = False,
    has_date: bool = False
) -> None:
    """테스트용 엑셀 파일 생성

    Args:
        filepath: 생성할 파일 경로
        sheet_name: 시트 이름 (기본: "summary")
        header: 1행 헤더 데이터
        data_row: 2행 데이터
        has_merged_cells: 1행에 병합 셀 포함 여부
        has_formula: 1행에 수식 포함 여부
        has_date: 1행에 날짜 포함 여부
    """
    workbook = Workbook()
    sheet = workbook.active

    if sheet is None:
        raise Exception("Worksheet creation failed")

    sheet.title = sheet_name

    # 기본 헤더와 데이터
    if header is None:
        header = ["Name", "Age", "City", "Score", "Status"]
    if data_row is None:
        data_row = ["John", 25, "Seoul", 95, "Active"]

    # 1행 (헤더) 작성
    if has_formula:
        # 수식 포함
        header[0] = "=SUM(A2:A10)"
    elif has_date:
        # 날짜 포함
        header[0] = datetime(2025, 1, 1)

    for col_idx, value in enumerate(header, start=1):
        sheet.cell(row=1, column=col_idx, value=value)

    # 병합 셀 추가
    if has_merged_cells:
        sheet.merge_cells('A1:B1')

    # 2행 (데이터) 작성
    if data_row:
        for col_idx, value in enumerate(data_row, start=1):
            sheet.cell(row=2, column=col_idx, value=value)

    # 파일 저장
    workbook.save(filepath)
    workbook.close()


def create_empty_file(filepath: str) -> None:
    """빈 엑셀 파일 생성 (시트는 있지만 2행 데이터가 비어있음)

    Args:
        filepath: 생성할 파일 경로
    """
    workbook = Workbook()
    sheet = workbook.active
    if sheet:
        sheet.title = "summary"
        # 헤더는 있지만 2행은 비어있음
        sheet.cell(row=1, column=1, value="Header1")
        sheet.cell(row=1, column=2, value="Header2")
        # 2행에 빈 값을 명시적으로 설정
        sheet.cell(row=2, column=1, value=None)
        sheet.cell(row=2, column=2, value=None)
    workbook.save(filepath)
    workbook.close()


def create_file_without_summary(filepath: str) -> None:
    """summary 시트가 없는 엑셀 파일 생성

    Args:
        filepath: 생성할 파일 경로
    """
    workbook = Workbook()
    sheet = workbook.active
    if sheet:
        sheet.title = "Sheet1"
        sheet.cell(row=1, column=1, value="Test")
    workbook.save(filepath)
    workbook.close()


def create_multiple_test_files(
    count: int,
    directory: str,
    prefix: str = "test_file"
) -> List[str]:
    """여러 개의 테스트 엑셀 파일 생성

    Args:
        count: 생성할 파일 개수
        directory: 저장 디렉토리
        prefix: 파일명 접두사

    Returns:
        생성된 파일 경로 리스트
    """
    file_paths = []

    for i in range(count):
        filepath = os.path.join(directory, f"{prefix}_{i+1}.xlsx")
        header = ["Name", "Age", "City", "Score", "Status"]
        data_row = [f"Person{i+1}", 20 + i, f"City{i+1}", 80 + i, "Active"]

        create_test_excel_file(
            filepath=filepath,
            header=header,
            data_row=data_row
        )
        file_paths.append(filepath)

    return file_paths


def create_files_with_different_columns(
    directory: str
) -> List[str]:
    """열 개수가 다른 파일들 생성

    Args:
        directory: 저장 디렉토리

    Returns:
        생성된 파일 경로 리스트
    """
    file_paths = []

    # 파일 1: 3개 열
    filepath1 = os.path.join(directory, "file_3_columns.xlsx")
    create_test_excel_file(
        filepath=filepath1,
        header=["A", "B", "C"],
        data_row=[1, 2, 3]
    )
    file_paths.append(filepath1)

    # 파일 2: 5개 열
    filepath2 = os.path.join(directory, "file_5_columns.xlsx")
    create_test_excel_file(
        filepath=filepath2,
        header=["A", "B", "C", "D", "E"],
        data_row=[1, 2, 3, 4, 5]
    )
    file_paths.append(filepath2)

    # 파일 3: 10개 열
    filepath3 = os.path.join(directory, "file_10_columns.xlsx")
    create_test_excel_file(
        filepath=filepath3,
        header=["A", "B", "C", "D", "E", "F", "G", "H", "I", "J"],
        data_row=[1, 2, 3, 4, 5, 6, 7, 8, 9, 10]
    )
    file_paths.append(filepath3)

    return file_paths


def get_temp_directory() -> str:
    """임시 디렉토리 생성 및 경로 반환

    Returns:
        임시 디렉토리 경로
    """
    return tempfile.mkdtemp(prefix="merger_test_")


def cleanup_test_files(directory: str) -> None:
    """테스트 파일 정리

    Args:
        directory: 정리할 디렉토리
    """
    import shutil
    if os.path.exists(directory):
        shutil.rmtree(directory)
