"""엑셀 파일 읽기 및 데이터 추출 (Phase 1, Phase 2 강화)

Single Responsibility: 엑셀 파일에서 데이터를 읽고 추출하는 책임만 담당
"""
import os
from typing import Any, List, Optional, Tuple
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook.workbook import Workbook
from src.logger import Logger


class ExcelReader:
    """엑셀 파일 읽기 및 데이터 추출 클래스"""

    SUMMARY_SHEET_NAME = "summary"

    def __init__(self, logger: Optional[Logger] = None) -> None:
        """ExcelReader 초기화

        Args:
            logger: Logger 인스턴스 (None이면 로깅 안 함)
        """
        self.logger = logger

    def read_file(self, file_path: str) -> Tuple[Optional[List[Any]], Optional[List[Any]], str]:  # noqa: E501
        """엑셀 파일에서 헤더(1행)와 데이터(2행) 읽기

        Args:
            file_path: 읽을 엑셀 파일 경로

        Returns:
            (헤더 데이터, 2행 데이터, 상태 메시지) 튜플
            - 성공: (header, row2, "SUCCESS")
            - summary 시트 없음: (None, None, "NO_SUMMARY_SHEET")
            - 2행 비어있음: (None, None, "ROW2_EMPTY")
            - 파일 잠김: (None, None, "FILE_LOCKED")
            - 파일 읽기 오류: (None, None, "FILE_ERROR")
        """
        file_name = os.path.basename(file_path)

        try:
            workbook = load_workbook(file_path, data_only=True)
        except PermissionError:
            if self.logger:
                self.logger.error("파일이 잠겨있습니다. 다른 프로그램에서 파일을 닫아주세요.", file_name)
            return (None, None, "FILE_LOCKED")
        except Exception as e:
            if self.logger:
                self.logger.error(f"파일 읽기 실패: {e}", file_name)
            return (None, None, "FILE_ERROR")

        # summary 시트 찾기 (대소문자 구분 없이)
        summary_sheet = self._find_summary_sheet(workbook)

        if summary_sheet is None:
            if self.logger:
                self.logger.skip("summary 시트가 없습니다.", file_name)
            workbook.close()
            return (None, None, "NO_SUMMARY_SHEET")

        # 1행(헤더) 추출
        header = self._extract_row(summary_sheet, 1)

        # 2행 추출
        row2_data = self._extract_row(summary_sheet, 2)

        workbook.close()

        # 2행 전체가 비어있는지 확인
        if self._is_row_empty(row2_data):
            if self.logger:
                self.logger.skip("summary 시트 2행 전체가 비어있습니다.", file_name)
            return (None, None, "ROW2_EMPTY")

        if self.logger:
            self.logger.success(f"{len(row2_data)} columns extracted", file_name)
        return (header, row2_data, "SUCCESS")

    def _find_summary_sheet(self, workbook: Workbook) -> Optional[Worksheet]:
        """summary 시트 찾기 (대소문자 구분 없이)

        Args:
            workbook: openpyxl Workbook 객체

        Returns:
            summary 시트 Worksheet 객체, 없으면 None
        """
        for sheet_name in workbook.sheetnames:
            if sheet_name.lower() == self.SUMMARY_SHEET_NAME.lower():
                return workbook[sheet_name]
        return None

    def _extract_row(self, sheet: Worksheet, row_number: int) -> List[Any]:
        """시트에서 특정 행의 데이터 추출

        Args:
            sheet: Worksheet 객체
            row_number: 추출할 행 번호 (1-based)

        Returns:
            행 데이터 리스트
        """
        row_data = []
        for cell in sheet[row_number]:
            row_data.append(cell.value)

        # 뒤에서부터 None 값 제거 (빈 컬럼 제거)
        while row_data and row_data[-1] is None:
            row_data.pop()

        return row_data

    def _is_row_empty(self, row_data: List[Any]) -> bool:
        """행 데이터가 전체 비어있는지 확인

        Args:
            row_data: 행 데이터 리스트

        Returns:
            전체 비어있으면 True
        """
        if not row_data:
            return True

        # 모든 셀이 None이거나 빈 문자열인지 확인
        for cell_value in row_data:
            if cell_value is not None and str(cell_value).strip() != "":
                return False

        return True
